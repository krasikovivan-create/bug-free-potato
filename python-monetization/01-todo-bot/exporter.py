"""Экспорт задач в Excel (.xlsx) через openpyxl."""
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from db import Task

HEADERS = ["№", "Task", "Status", "Created", "Reminder", "Done at"]
DATE_FORMAT = "yyyy-mm-dd hh:mm"


def _naive_local(dt: datetime | None, zone: ZoneInfo) -> datetime | None:
    """Excel не умеет часовые пояса: переводим в время пользователя и убираем tzinfo."""
    return dt.astimezone(zone).replace(tzinfo=None) if dt else None


def build_xlsx(tasks: list[Task], zone: ZoneInfo) -> bytes:
    """Собирает Excel-файл со всеми задачами и возвращает его содержимое (bytes)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tasks"
    ws.append(HEADERS)

    for task in tasks:
        ws.append([
            task.num,
            None,  # текст запишем ниже отдельно — см. комментарий
            "Done" if task.done else "Open",
            _naive_local(task.created_at, zone),
            _naive_local(task.remind_at, zone),
            _naive_local(task.done_at, zone),
        ])
        # Важно: openpyxl считает строку, начинающуюся с «=», формулой. Пользователь может
        # прислать боту «=HYPERLINK(...)», и Excel его выполнит. Поэтому тип ячейки ставим
        # явно — «строка», и чистим управляющие символы, на которых openpyxl падает.
        cell = ws.cell(row=ws.max_row, column=2)
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", task.text)
        cell.data_type = "s"
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # --- оформление ---
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = DATE_FORMAT
    for letter, width in zip("ABCDEF", (6, 60, 10, 18, 18, 18)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"  # шапка не уезжает при прокрутке
    ws.auto_filter.ref = ws.dimensions  # фильтры в шапке

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

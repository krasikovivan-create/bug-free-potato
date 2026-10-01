from datetime import datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo
from zipfile import ZipFile

from openpyxl import load_workbook

from db import Database
from exporter import HEADERS, build_xlsx

NOW = datetime(2026, 10, 1, 12, 0, tzinfo=timezone.utc)


def make_db(tmp_path):
    return Database(str(tmp_path / "t.db"))


def load(data):
    return load_workbook(BytesIO(data)).active


def test_export_contains_all_tasks(tmp_path):
    db = make_db(tmp_path)
    db.add_task(1, "Buy milk", now=NOW)
    db.add_task(1, "Call Bob", now=NOW)
    db.mark_done(1, 1, now=NOW)
    db.set_reminder(1, 2, datetime(2026, 10, 2, 9, 0, tzinfo=timezone.utc))

    ws = load(build_xlsx(db.list_tasks(1, include_done=True), ZoneInfo("Europe/Berlin")))
    rows = list(ws.iter_rows(values_only=True))

    assert list(rows[0]) == HEADERS
    open_task, done_task = rows[1], rows[2]  # открытые идут первыми
    assert open_task[:3] == (2, "Call Bob", "Open")
    assert done_task[:3] == (1, "Buy milk", "Done")
    # время переведено в пояс пользователя (Берлин, UTC+2)
    assert open_task[3] == datetime(2026, 10, 1, 14, 0)
    assert open_task[4] == datetime(2026, 10, 2, 11, 0)
    assert open_task[5] is None
    assert done_task[5] == datetime(2026, 10, 1, 14, 0)


def test_formula_text_is_stored_as_plain_string(tmp_path):
    db = make_db(tmp_path)
    for text in ['=HYPERLINK("http://evil.example","click")', "=1+1", "+cmd", "-5", "@SUM(A1)"]:
        db.add_task(1, text, now=NOW)

    data = build_xlsx(db.list_tasks(1), ZoneInfo("UTC"))

    with ZipFile(BytesIO(data)) as z:
        sheet_xml = z.read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in sheet_xml  # ни одной формулы в файле

    ws = load(data)
    assert [row[1] for row in ws.iter_rows(min_row=2, values_only=True)] == [
        '=HYPERLINK("http://evil.example","click")', "=1+1", "+cmd", "-5", "@SUM(A1)",
    ]


def test_control_characters_do_not_break_export(tmp_path):
    db = make_db(tmp_path)
    db.add_task(1, "bad\x00\x07char", now=NOW)
    ws = load(build_xlsx(db.list_tasks(1), ZoneInfo("UTC")))
    assert ws["B2"].value == "badchar"


def test_unicode_survives(tmp_path):
    db = make_db(tmp_path)
    db.add_task(1, "Купить молоко 🥛", now=NOW)
    ws = load(build_xlsx(db.list_tasks(1), ZoneInfo("UTC")))
    assert ws["B2"].value == "Купить молоко 🥛"
